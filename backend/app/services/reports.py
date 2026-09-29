from __future__ import annotations

import csv
import io
import time
import uuid
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import sqlalchemy as sa
from fastapi import Request
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.arm112 import FIELD_LABELS
from app.core.config import settings
from app.core.exceptions import BusinessRuleError, NotFoundError
from app.core.logging import get_logger
from app.core.pagination import PageParams
from app.core.security import new_opaque_token, utcnow
from app.models.card import IncidentCard
from app.models.enums import (
    LESSON_PURPOSE_TITLES,
    AuditAction,
    LessonPurpose,
    ReportFormat,
    ReportStatus,
    ReportType,
)
from app.models.grading import ErrorRecord, Evaluation
from app.models.report import Certificate, Report
from app.models.training import CardAttempt, Lesson, LessonParticipant
from app.models.user import User
from app.repositories.grading import EvaluationRepository
from app.repositories.system import CertificateRepository, ReportRepository
from app.repositories.training import AttemptRepository, LessonRepository
from app.repositories.users import UserRepository
from app.schemas.report import CertificateIssueRequest, ReportCreate
from app.services.analytics import AnalyticsService
from app.services.audit import AuditService

logger = get_logger(__name__)


class ReportService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.reports = ReportRepository(session)
        self.certificates = CertificateRepository(session)
        self.lessons = LessonRepository(session)
        self.attempts = AttemptRepository(session)
        self.evaluations = EvaluationRepository(session)
        self.users = UserRepository(session)
        self.analytics = AnalyticsService(session)
        self.audit = AuditService(session)
        settings.ensure_storage()

    # ------------------------------------------------------------------ отчёты
    async def list_reports(
        self, actor: User, params: PageParams, *, own_only: bool = False
    ) -> tuple[Sequence[Report], int]:
        conditions = []
        if own_only:
            conditions.append(sa.or_(Report.student_id == actor.id, Report.requested_by_id == actor.id))
        return await self.reports.paginate(params, *conditions)

    async def get(self, report_id: uuid.UUID) -> Report:
        report = await self.reports.get(report_id)
        if report is None:
            raise NotFoundError("Отчёт не найден")
        return report

    async def create(
        self, data: ReportCreate, actor: User, request: Request | None = None
    ) -> Report:
        started = time.perf_counter()
        report = Report(
            type=data.type,
            format=data.format,
            status=ReportStatus.PENDING,
            title=await self._title(data),
            requested_by_id=actor.id,
            lesson_id=data.lesson_id,
            student_id=data.student_id,
            group_id=data.group_id,
            params={
                **data.params,
                "date_from": data.date_from.isoformat() if data.date_from else None,
                "date_to": data.date_to.isoformat() if data.date_to else None,
            },
        )
        self.session.add(report)
        await self.session.flush()

        try:
            payload, rows = await self._collect(data)
            report.data = payload
            report.row_count = len(rows)
            if data.format is not ReportFormat.JSON:
                path = await self._render_file(report, payload, rows)
                report.file_path = str(path)
                report.file_size = path.stat().st_size
            report.status = ReportStatus.READY
            report.generated_at = utcnow()
        except Exception as exc:  # noqa: BLE001 — отчёт не должен ронять запрос
            report.status = ReportStatus.FAILED
            report.error = str(exc)[:1000]
            logger.exception("report_failed")
        finally:
            report.generation_ms = int((time.perf_counter() - started) * 1000)
            await self.session.flush()

        if report.generation_ms > settings.REPORT_TIMEOUT_SECONDS * 1000:
            logger.warning(
                "report_slow",
                extra={"report_id": str(report.id), "generation_ms": report.generation_ms},
            )

        await self.audit.log(
            AuditAction.EXPORT,
            actor=actor,
            object_type="report",
            object_id=report.id,
            summary=f"Сформирован отчёт «{report.title}» ({report.format.value})",
            after={"status": report.status.value, "generation_ms": report.generation_ms},
            request=request,
        )
        return report

    async def file_path(self, report_id: uuid.UUID) -> Path:
        report = await self.get(report_id)
        if not report.file_path:
            raise BusinessRuleError("У отчёта нет файла: он сформирован в формате JSON")
        path = Path(report.file_path)
        if not path.exists():
            raise NotFoundError("Файл отчёта не найден на диске")
        return path

    # --------------------------------------------------------------- сбор данных
    async def _collect(self, data: ReportCreate) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        if data.type is ReportType.LESSON:
            if not data.lesson_id:
                raise BusinessRuleError("Для отчёта по занятию нужен lesson_id")
            return await self._lesson_report(data.lesson_id)
        if data.type is ReportType.PROGRESS:
            if not data.student_id:
                raise BusinessRuleError("Для отчёта по успеваемости нужен student_id")
            return await self._progress_report(data.student_id)
        if data.type is ReportType.ATTESTATION:
            if not data.lesson_id:
                raise BusinessRuleError("Для протокола аттестации нужен lesson_id")
            return await self._attestation_report(data.lesson_id)
        if data.type in (ReportType.GROUP_STATS, ReportType.ERRORS, ReportType.TIMING):
            summary = await self.analytics.summary(
                lesson_id=data.lesson_id,
                student_id=data.student_id,
                group_id=data.group_id,
                date_from=data.date_from,
                date_to=data.date_to,
            )
            rows_key = {
                ReportType.GROUP_STATS: "score_dynamics",
                ReportType.ERRORS: "top_errors",
                ReportType.TIMING: "timing",
            }[data.type]
            return summary, list(summary.get(rows_key) or [])
        raise BusinessRuleError(f"Тип отчёта {data.type.value} не поддерживается")

    async def _lesson_report(self, lesson_id: uuid.UUID) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        """Отчёт о практическом занятии: действия, замечания, время, отклонение, грамматика."""
        lesson = await self.lessons.get_full(lesson_id)
        if lesson is None:
            raise NotFoundError("Занятие не найдено")

        stmt = (
            sa.select(CardAttempt, Evaluation, IncidentCard, User)
            .join(Evaluation, Evaluation.attempt_id == CardAttempt.id, isouter=True)
            .join(IncidentCard, IncidentCard.id == CardAttempt.card_id)
            .join(User, User.id == CardAttempt.student_id)
            .where(CardAttempt.lesson_id == lesson_id)
            .order_by(User.full_name, CardAttempt.sequence_no)
        )
        rows: list[dict[str, Any]] = []
        for attempt, evaluation, card, student in (await self.session.execute(stmt)).all():
            errors = []
            if evaluation is not None:
                error_rows = (
                    await self.session.execute(
                        sa.select(ErrorRecord).where(ErrorRecord.evaluation_id == evaluation.id)
                    )
                ).scalars().all()
                errors = [
                    {
                        "category": e.category.value,
                        "severity": e.severity.value,
                        "message": e.message,
                        "field": e.field_code,
                    }
                    for e in error_rows
                ]
            rows.append(
                {
                    "student": student.full_name,
                    "card_no": card.card_no,
                    "sequence_no": attempt.sequence_no,
                    "status": attempt.status.value,
                    "duration_seconds": round((attempt.duration_ms or 0) / 1000, 2),
                    "norm_seconds": attempt.norm_seconds,
                    "time_delta_seconds": attempt.time_delta_seconds,
                    "is_overtime": attempt.is_overtime,
                    "score": evaluation.score if evaluation else None,
                    "passed": evaluation.passed if evaluation else None,
                    "grammar_score": evaluation.grammar_score if evaluation else None,
                    "error_count": evaluation.error_count if evaluation else 0,
                    "errors": errors,
                    "teacher_comment": evaluation.teacher_comment if evaluation else None,
                }
            )

        summary = await self.analytics.summary(lesson_id=lesson_id)
        payload = {
            "lesson": {
                "id": str(lesson.id),
                "title": lesson.title,
                "mode": lesson.mode.value,
                "status": lesson.status.value,
                "started_at": lesson.started_at.isoformat() if lesson.started_at else None,
                "finished_at": lesson.finished_at.isoformat() if lesson.finished_at else None,
                "time_limit_seconds": lesson.time_limit_seconds,
                "participants": len(lesson.participants),
            },
            "summary": summary,
            "rows": rows,
        }
        return payload, rows

    async def _progress_report(self, student_id: uuid.UUID) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        student = await self.users.get_or_fail(student_id, "Обучающийся не найден")
        summary = await self.analytics.summary(student_id=student_id)
        history = await self.analytics.student_progress(student_id)
        rows = [
            {
                "recorded_at": item.recorded_at.isoformat(),
                "score": item.score,
                "passed": item.passed,
                "duration_seconds": round((item.duration_ms or 0) / 1000, 2),
                "error_count": item.error_count,
            }
            for item in history
        ]
        forecast: dict[str, Any] | None
        try:
            forecast = await self.analytics.forecast(student_id)
            forecast = {k: (str(v) if isinstance(v, uuid.UUID) else v) for k, v in forecast.items()}
        except NotFoundError:
            forecast = None

        payload = {
            "student": {"id": str(student.id), "full_name": student.full_name},
            "summary": summary,
            "forecast": forecast,
            "rows": rows,
        }
        return payload, rows

    async def _attestation_report(
        self, lesson_id: uuid.UUID
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        lesson = await self.lessons.get_full(lesson_id)
        if lesson is None:
            raise NotFoundError("Занятие не найдено")

        teacher = await self.users.get(lesson.teacher_id)
        threshold = lesson.passing_score or settings.ATTESTATION_PASS_SCORE
        rows: list[dict[str, Any]] = []
        for participant in lesson.participants:
            student = await self.users.get(participant.student_id)
            errors_total = int(
                (
                    await self.session.execute(
                        sa.select(sa.func.count())
                        .select_from(ErrorRecord)
                        .join(Evaluation, Evaluation.id == ErrorRecord.evaluation_id)
                        .where(
                            Evaluation.lesson_id == lesson_id,
                            Evaluation.student_id == participant.student_id,
                        )
                    )
                ).scalar_one()
            )
            rows.append(
                {
                    "student": student.full_name if student else str(participant.student_id),
                    "cards_issued": participant.cards_issued,
                    "cards_submitted": participant.cards_submitted,
                    "score": participant.final_score,
                    "errors": errors_total,
                    "verdict": _verdict(participant.is_passed),
                }
            )

        payload = {
            "lesson": {
                "id": str(lesson.id),
                "title": lesson.title,
                "purpose": LESSON_PURPOSE_TITLES[lesson.purpose],
                "mode": lesson.mode.value,
                "finished_at": lesson.finished_at.isoformat() if lesson.finished_at else None,
            },
            "teacher": teacher.full_name if teacher else None,
            "passing_score": float(threshold),
            "participants_total": len(rows),
            "passed_total": sum(1 for row in rows if row["verdict"] == "зачёт"),
            "rows": rows,
        }
        return payload, rows

    # ---------------------------------------------------------------- рендеринг
    async def _render_file(self, report: Report, payload: dict[str, Any], rows: list[dict]) -> Path:
        stamp = utcnow().strftime("%Y%m%d_%H%M%S")
        name = f"{report.type.value}_{stamp}_{report.id.hex[:8]}"
        if report.format is ReportFormat.CSV:
            return _write_csv(settings.reports_path / f"{name}.csv", rows)
        if report.format is ReportFormat.PDF:
            return _write_pdf(settings.reports_path / f"{name}.pdf", report.title, payload, rows)
        if report.format is ReportFormat.XLSX:
            return _write_xlsx(settings.reports_path / f"{name}.xlsx", report.title, payload, rows)
        if report.format is ReportFormat.XML:
            return _write_xml(settings.reports_path / f"{name}.xml", report, payload, rows)
        raise BusinessRuleError(f"Формат {report.format.value} не поддерживается")

    async def _title(self, data: ReportCreate) -> str:
        titles = {
            ReportType.LESSON: "Отчёт о практическом занятии",
            ReportType.PROGRESS: "Отчёт об успеваемости обучающегося",
            ReportType.GROUP_STATS: "Статистика обучения",
            ReportType.ERRORS: "Отчёт по типичным ошибкам",
            ReportType.TIMING: "Отчёт по соблюдению нормативов времени",
            ReportType.AUDIT: "Выгрузка журнала аудита",
            ReportType.ATTESTATION: "Протокол аттестационного мероприятия",
        }
        return titles.get(data.type, "Отчёт")

    # ------------------------------------------------------------- сертификаты
    async def issue_certificate(
        self, data: CertificateIssueRequest, actor: User, request: Request | None = None
    ) -> Certificate:
        student = await self.users.get_or_fail(data.student_id, "Обучающийся не найден")
        avg_score = await self.evaluations.avg_score(Evaluation.student_id == student.id)
        if avg_score is None:
            raise BusinessRuleError("У обучающегося нет оценённых занятий — сертификат не выдаётся")

        attestation = await self._attestation_result(student.id, data.lesson_id)
        if attestation is not None:
            score, passed = attestation
            if not passed:
                raise BusinessRuleError(
                    f"Аттестация не пройдена (результат {score:.1f} из 100) — сертификат не выдаётся",
                    code="attestation_not_passed",
                )
            avg_score = score

        certificate = Certificate(
            serial=await self.certificates.next_serial(),
            student_id=student.id,
            lesson_id=data.lesson_id,
            program_name=data.program_name,
            issued_by_id=actor.id,
            score=round(avg_score, 2),
            hours=data.hours,
            verification_code=new_opaque_token(12),
        )
        self.session.add(certificate)
        await self.session.flush()

        path = _write_certificate_pdf(
            settings.certificates_path / f"certificate_{certificate.serial.replace('/', '_')}.pdf",
            certificate,
            student,
        )
        certificate.file_path = str(path)
        await self.session.flush()

        await self.audit.log(
            AuditAction.EXPORT,
            actor=actor,
            object_type="certificate",
            object_id=certificate.id,
            summary=f"Выдан сертификат {certificate.serial} — {student.full_name}",
            request=request,
        )
        return certificate

    async def _attestation_result(
        self, student_id: uuid.UUID, lesson_id: uuid.UUID | None
    ) -> tuple[float, bool] | None:
        stmt = (
            sa.select(LessonParticipant.final_score, LessonParticipant.is_passed)
            .join(Lesson, Lesson.id == LessonParticipant.lesson_id)
            .where(
                LessonParticipant.student_id == student_id,
                LessonParticipant.is_passed.is_not(None),
                Lesson.purpose.in_((LessonPurpose.ATTESTATION, LessonPurpose.REFRESHER)),
            )
            .order_by(Lesson.finished_at.desc().nullslast())
        )
        if lesson_id:
            stmt = stmt.where(Lesson.id == lesson_id)
        row = (await self.session.execute(stmt.limit(1))).first()
        if row is None:
            if lesson_id is not None:
                lesson = await self.lessons.get(lesson_id)
                if lesson is not None and lesson.purpose in (
                    LessonPurpose.ATTESTATION,
                    LessonPurpose.REFRESHER,
                ):
                    raise BusinessRuleError(
                        "Итоги аттестации не подведены: завершите занятие, затем выдавайте сертификат",
                        code="attestation_not_finished",
                    )
            return None
        return float(row[0] or 0.0), bool(row[1])

    async def list_certificates(
        self, params: PageParams, *, student_id: uuid.UUID | None = None
    ) -> tuple[Sequence[Certificate], int]:
        conditions = []
        if student_id:
            conditions.append(Certificate.student_id == student_id)
        return await self.certificates.paginate(params, *conditions)

    async def certificate_file(self, certificate_id: uuid.UUID, actor: User) -> tuple[Certificate, Path]:
        certificate = await self.certificates.get_or_fail(certificate_id, "Сертификат не найден")
        if certificate.student_id != actor.id and not actor.has_permission("certificates:issue"):
            from app.core.exceptions import PermissionDeniedError

            raise PermissionDeniedError("Сертификат другого обучающегося недоступен")
        if not certificate.file_path or not Path(certificate.file_path).exists():
            raise NotFoundError("Файл сертификата не найден")
        return certificate, Path(certificate.file_path)


def _verdict(is_passed: bool | None) -> str:
    if is_passed is None:
        return "не оценивался"
    return "зачёт" if is_passed else "незачёт"


# --------------------------------------------------------------------- рендеры
def _write_xlsx(path: Path, title: str, payload: dict[str, Any], rows: list[dict[str, Any]]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Данные"
    if rows:
        headers = list({key: None for row in rows for key in row})
        sheet.append([FIELD_LABELS.get(header, header) for header in headers])
        for row in rows:
            sheet.append([_flatten(row.get(header)) for header in headers])
        for column_index, header in enumerate(headers, start=1):
            width = max(len(str(header)), 12)
            for row in rows[:200]:
                width = max(width, min(60, len(str(_flatten(row.get(header)) or ""))))
            sheet.column_dimensions[get_column_letter(column_index)].width = width + 2
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
    else:
        sheet.append(["Нет данных"])

    summary = workbook.create_sheet("Сводка")
    summary.append(["Отчёт", title])
    for key, value in payload.items():
        if isinstance(value, dict | list):
            continue
        summary.append([FIELD_LABELS.get(key, key), _flatten(value)])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 48
    for cell in summary["A"]:
        cell.font = Font(bold=True)

    workbook.save(path)
    return path


def _write_xml(path: Path, report: Report, payload: dict[str, Any], rows: list[dict[str, Any]]) -> Path:
    from xml.etree import ElementTree as ET

    path.parent.mkdir(parents=True, exist_ok=True)
    root = ET.Element(
        "report",
        {
            "id": str(report.id),
            "type": report.type.value,
            "title": report.title,
            "generated_at": (report.generated_at or utcnow()).isoformat(),
        },
    )
    summary = ET.SubElement(root, "summary")
    for key, value in payload.items():
        if isinstance(value, dict | list):
            continue
        item = ET.SubElement(summary, "item", {"code": key, "label": FIELD_LABELS.get(key, key)})
        item.text = str(_flatten(value))

    rows_node = ET.SubElement(root, "rows", {"count": str(len(rows))})
    for row in rows:
        row_node = ET.SubElement(rows_node, "row")
        for key, value in row.items():
            field = ET.SubElement(
                row_node, "field", {"code": _xml_name(key), "label": FIELD_LABELS.get(key, key)}
            )
            field.text = "" if value is None else str(_flatten(value))

    ET.indent(root, space="  ")
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)
    return path


def _xml_name(value: str) -> str:
    """Код поля, безопасный для XML-атрибута."""
    return "".join(char if char.isalnum() or char in "_-." else "_" for char in value)


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    buffer = io.StringIO()
    if rows:
        fieldnames = list({key: None for row in rows for key in row})
        writer = csv.DictWriter(buffer, fieldnames=fieldnames, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: _flatten(v) for k, v in row.items()})
    else:
        buffer.write("Нет данных\n")
    # utf-8-sig — чтобы Excel корректно открывал русские заголовки.
    path.write_text(buffer.getvalue(), encoding="utf-8-sig")
    return path


def _flatten(value: Any) -> Any:
    if isinstance(value, dict | list):
        import json

        return json.dumps(value, ensure_ascii=False)
    return value


def _register_font() -> str:
    """Подбирает шрифт с кириллицей: reportlab по умолчанию её не умеет."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
        Path("/Library/Fonts/Arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                pdfmetrics.registerFont(TTFont("AppFont", str(candidate)))
                return "AppFont"
            except Exception:  # noqa: BLE001
                continue
    logger.warning("cyrillic_font_not_found", extra={"fallback": "Helvetica"})
    return "Helvetica"


def _write_pdf(path: Path, title: str, payload: dict[str, Any], rows: list[dict[str, Any]]) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    path.parent.mkdir(parents=True, exist_ok=True)
    font = _register_font()
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), title=title)
    heading = ParagraphStyle("Heading", fontName=font, fontSize=15, leading=19, spaceAfter=10)
    normal = ParagraphStyle("Normal", fontName=font, fontSize=8.5, leading=11)

    story: list[Any] = [Paragraph(title, heading)]
    lesson = payload.get("lesson") or payload.get("student") or {}
    meta = " · ".join(f"{key}: {value}" for key, value in lesson.items() if value is not None)
    if meta:
        story.append(Paragraph(meta, normal))
    summary = payload.get("summary") or {}
    if summary:
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                "Средний балл: {avg} · Попыток: {total} · С превышением норматива: {over}".format(
                    avg=summary.get("avg_score", "—"),
                    total=summary.get("attempts_total", 0),
                    over=summary.get("attempts_overtime", 0),
                ),
                normal,
            )
        )
    story.append(Spacer(1, 12))

    if rows:
        columns = [key for key in rows[0] if key != "errors"][:9]
        data = [[Paragraph(str(col), normal) for col in columns]]
        for row in rows[:400]:
            data.append([Paragraph(str(_flatten(row.get(col, ""))), normal) for col in columns])
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7BDC7")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF4")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTNAME", (0, 0), (-1, -1), font),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("Нет данных за выбранный период", normal))

    doc.build(story)
    return path


def _write_certificate_pdf(path: Path, certificate: Certificate, student: User) -> Path:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    path.parent.mkdir(parents=True, exist_ok=True)
    font = _register_font()
    doc = SimpleDocTemplate(
        str(path), pagesize=landscape(A4), title=f"Сертификат {certificate.serial}",
        leftMargin=25 * mm, rightMargin=25 * mm, topMargin=25 * mm, bottomMargin=20 * mm,
    )
    center_big = ParagraphStyle("Big", fontName=font, fontSize=26, leading=32, alignment=1)
    center = ParagraphStyle("Center", fontName=font, fontSize=13, leading=18, alignment=1)
    small = ParagraphStyle("Small", fontName=font, fontSize=9, leading=12, alignment=1)

    issued = certificate.issued_at or utcnow()
    story = [
        Paragraph("СЕРТИФИКАТ", center_big),
        Spacer(1, 6 * mm),
        Paragraph(f"№ {certificate.serial}", center),
        Spacer(1, 10 * mm),
        Paragraph("выдан", small),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>{student.full_name}</b>", center_big),
        Spacer(1, 8 * mm),
        Paragraph(f"в подтверждение прохождения программы<br/>«{certificate.program_name}»", center),
        Spacer(1, 8 * mm),
        Paragraph(
            f"Итоговый результат: {certificate.score if certificate.score is not None else '—'} из 100"
            + (f" · Объём: {certificate.hours} ч." if certificate.hours else ""),
            center,
        ),
        Spacer(1, 12 * mm),
        Paragraph(f"Дата выдачи: {issued:%d.%m.%Y}", small),
        Paragraph(f"Код проверки подлинности: {certificate.verification_code}", small),
        Spacer(1, 6 * mm),
        Paragraph("Учебный комплекс подготовки операторов ДДС города Москвы", small),
    ]
    doc.build(story)
    return path
