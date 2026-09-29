from __future__ import annotations

import uuid
from collections.abc import Sequence
from typing import Any

import sqlalchemy as sa
from fastapi import Request
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.arm112 import service_title
from app.core.exceptions import BusinessRuleError, NotFoundError
from app.core.logging import get_logger
from app.core.pagination import PageParams
from app.models.catalog import IncidentCategory, IncidentType
from app.models.enums import AuditAction
from app.models.user import User
from app.repositories.content import CategoryRepository, IncidentTypeRepository
from app.schemas.catalog import IncidentTypeImportItem
from app.services.audit import AuditService

logger = get_logger(__name__)

#: Варианты заголовков по убыванию приоритета: берётся колонка, совпавшая с более ранним
#: вариантом (в таблице заказчика «112 - Признак.1 (тип происшествия)» стоит левее
#: «Итоговый тип происшествия», но название типа — именно итоговый).
DEFAULT_COLUMN_MAP: dict[str, tuple[str, ...]] = {
    "code": ("код екп", "код происшествия", "код", "номер", "id"),
    "category_name": ("категория происшествия", "категория", "класс", "группа"),
    "object": ("что случилось", "объект", "признак"),
    "detail": ("объект уточнение", "уточнение", "подробность", "деталь"),
    "name": ("итоговый тип", "наименование", "тип происшествия", "происшествие"),
    "main_service": ("главная служба", "основная служба", "служба"),
}

AGENCY_PREFIXES = ("классификатор", "квс", "мчс", "мвд", "смп", "фсб", "жкх")

#: В скольких первых строках ищем заголовок таблицы.
HEADER_LOOKUP_ROWS = 25

#: Названия служб из таблицы → коды служб оповещения (app.core.arm112).
SERVICE_ALIASES: dict[str, str] = {
    "police": "102",
    "полиция": "102",
    "мвд": "102",
    "гибдд": "dps",
    "дпс": "dps",
    "fire": "101",
    "пожарная охрана": "101",
    "мчс": "101",
    "псг": "101",
    "medical": "103",
    "скорая": "103",
    "смп": "103",
    "скорая помощь": "103",
    "газ": "104",
    "мосгаз": "104",
    "водоканал": "mosvodokanal",
    "мосводоканал": "mosvodokanal",
    "моэк": "moek",
    "жилищник": "zhilishnik",
    "жкх": "zhilishnik",
}


class IncidentCatalogService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.types = IncidentTypeRepository(session)
        self.categories = CategoryRepository(session)
        self.audit = AuditService(session)

    # ------------------------------------------------------------------ чтение
    async def list_types(
        self,
        params: PageParams,
        *,
        category_name: str | None = None,
        category_id: uuid.UUID | None = None,
        service: str | None = None,
        query: str | None = None,
    ) -> tuple[Sequence[IncidentType], int]:
        conditions = [IncidentType.is_active.is_(True)]
        if category_name:
            conditions.append(IncidentType.category_name == category_name)
        if category_id:
            conditions.append(IncidentType.category_id == category_id)
        if service:
            conditions.append(IncidentType.main_service == service)
        if query:
            pattern = f"%{query.lower()}%"
            conditions.append(
                sa.or_(
                    sa.func.lower(IncidentType.name).like(pattern),
                    sa.func.lower(IncidentType.code).like(pattern),
                )
            )
        return await self.types.paginate(params, *conditions)

    async def get_type(self, key: str) -> IncidentType:
        """Тип происшествия по коду классификатора либо по идентификатору записи."""
        found = await self.types.find_one(IncidentType.code == key)
        if found is None:
            try:
                found = await self.types.get(uuid.UUID(key))
            except ValueError:
                found = None
        if found is None:
            raise NotFoundError(f"Тип происшествия «{key}» не найден в классификаторе")
        return found

    async def features_of(self, key: str) -> list[dict[str, Any]]:
        return list((await self.get_type(key)).features or [])

    async def categories_summary(self) -> list[dict[str, Any]]:
        """Категории классификатора со счётчиком типов — для выпадающих списков."""
        stmt = (
            sa.select(
                IncidentType.category_name,
                IncidentType.category_id,
                sa.func.count().label("types_total"),
            )
            .where(IncidentType.is_active.is_(True), IncidentType.category_name.is_not(None))
            .group_by(IncidentType.category_name, IncidentType.category_id)
            .order_by(IncidentType.category_name)
        )
        rows = (await self.session.execute(stmt)).all()

        codes: dict[uuid.UUID, str] = {}
        category_ids = [row.category_id for row in rows if row.category_id]
        if category_ids:
            for category in await self.categories.by_ids(category_ids):
                codes[category.id] = category.code

        return [
            {
                "category_name": row.category_name,
                "types_total": row.types_total,
                "category_id": row.category_id,
                "category_code": codes.get(row.category_id) if row.category_id else None,
            }
            for row in rows
        ]

    # ------------------------------------------------------------------ импорт
    async def import_types(
        self,
        items: Sequence[IncidentTypeImportItem],
        *,
        update_existing: bool = True,
        actor: User | None = None,
        request: Request | None = None,
    ) -> dict[str, Any]:
        """Сохраняет нормализованные записи классификатора."""
        category_ids = await self._category_index()
        imported = updated = skipped = 0
        warnings: list[str] = []

        for item in items:
            existing = await self.types.find_one(IncidentType.code == item.code)
            if existing is not None and not update_existing:
                skipped += 1
                continue

            values = {
                "name": item.name,
                "category_name": item.category_name,
                "category_id": category_ids.get((item.category_code or "").lower()),
                "features": [feature.model_dump() for feature in item.features],
                "main_service": item.main_service,
                "services": self._services_of(item),
                "agency_classifiers": item.agency_classifiers,
                "raw": item.raw,
            }
            if item.category_code and values["category_id"] is None:
                warnings.append(f"{item.code}: категория «{item.category_code}» не найдена")

            if existing is None:
                await self.types.create(code=item.code, **values)
                imported += 1
            else:
                await self.types.update(existing, **values)
                updated += 1

        await self.audit.log(
            AuditAction.CREATE,
            actor=actor,
            actor_username=None if actor else "система",
            object_type="incident_type",
            summary=(
                f"Импорт классификатора происшествий: добавлено {imported}, "
                f"обновлено {updated}, пропущено {skipped}"
            ),
            after={"imported": imported, "updated": updated, "skipped": skipped},
            request=request,
        )
        return {
            "imported": imported,
            "updated": updated,
            "skipped": skipped,
            "warnings": warnings[:50],
        }

    def parse_xlsx(
        self, payload: bytes, *, sheet: str | None = None, limit: int | None = None
    ) -> list[IncidentTypeImportItem]:
        import io

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]

        rows = worksheet.iter_rows(values_only=True)
        preamble = [row for _, row in zip(range(HEADER_LOOKUP_ROWS), rows, strict=False)]
        header, header_index, mapping = self._find_header(preamble)
        if header is None:
            raise BusinessRuleError(
                "Не удалось определить строку заголовков таблицы классификатора. "
                "Ожидаются колонки «Категория», «Тип происшествия», «Главная служба». "
                "Укажите лист параметром sheet либо передайте нормализованные данные JSON."
            )
        if header_index > 0:
            #: Шапка в две строки (у заказчика «Главная служба» стоит строкой выше основной
            #: шапки): пустые ячейки заголовка берём из строки над ним.
            above = preamble[header_index - 1] or ()
            header = tuple(
                cell if cell not in (None, "") else (above[i] if i < len(above) else None)
                for i, cell in enumerate(header)
            )
            mapping = self._map_columns(header)

        items: list[IncidentTypeImportItem] = []
        data_rows = [*preamble[header_index + 1 :], *rows]
        for row_no, row in enumerate(data_rows, start=header_index + 2):
            item = self._row_to_item(row, header, mapping, row_no)
            if item is not None:
                items.append(item)
            if limit and len(items) >= limit:
                break

        workbook.close()
        if not items:
            raise BusinessRuleError("В таблице не нашлось ни одной заполненной строки классификатора")
        return items

    # ---------------------------------------------------------------- приватное
    async def _category_index(self) -> dict[str, uuid.UUID]:
        stmt = sa.select(IncidentCategory.code, IncidentCategory.id)
        return {code.lower(): category_id for code, category_id in (await self.session.execute(stmt)).all()}

    def _services_of(self, item: IncidentTypeImportItem) -> list[dict[str, Any]]:
        """Службы классификатора приводятся к кодам списка оповещения."""
        result: list[dict[str, Any]] = []
        seen: set[str] = set()
        candidates = [(item.main_service, True), *[(name, False) for name in item.services]]
        for name, is_primary in candidates:
            code = normalize_service(name)
            if code is None or code in seen:
                continue
            seen.add(code)
            result.append(
                {
                    "code": code,
                    "name": service_title(code),
                    "is_primary": is_primary,
                    "reason": "ЕКП",
                }
            )
        return result

    def _find_header(self, rows: list[tuple]) -> tuple[tuple | None, int, dict[str, int]]:
        best: tuple[tuple | None, int, dict[str, int]] = (None, 0, {})
        for index, row in enumerate(rows):
            if not row:
                continue
            mapping = self._map_columns(row)
            if len(mapping) < 2:
                continue
            if "name" not in mapping and "category_name" not in mapping:
                continue
            if len(mapping) > len(best[2]):
                best = (row, index, mapping)
        return best

    def _map_columns(self, header: tuple) -> dict[str, int]:
        titles = [str(cell or "").strip().lower() for cell in header]
        mapping: dict[str, int] = {}
        used: set[int] = set()
        for field, variants in DEFAULT_COLUMN_MAP.items():
            best: tuple[int, int] | None = None   # (ранг варианта: точное совпадение раньше, индекс колонки)
            for index, title in enumerate(titles):
                if not title or title.startswith("unnamed") or index in used:
                    continue
                for rank, variant in enumerate(variants):
                    score = rank * 2 + (0 if title == variant else 1)
                    if (variant == title or variant in title) and (best is None or score < best[0]):
                        best = (score, index)
            if best is not None:
                mapping[field] = best[1]
                used.add(best[1])
        return mapping

    def _row_to_item(
        self, row: tuple, header: tuple, mapping: dict[str, int], row_no: int
    ) -> IncidentTypeImportItem | None:
        def value(field: str) -> str | None:
            index = mapping.get(field)
            if index is None or index >= len(row):
                return None
            cell = row[index]
            text = str(cell).strip() if cell is not None else ""
            return text or None

        name = value("name")
        category_name = value("category_name")
        if not name and not category_name:
            return None  # строка-разделитель

        features = [
            {"type": kind, "value": text}
            for kind, text in (("object", value("object")), ("detail", value("detail")))
            if text
        ]
        agency: dict[str, str] = {}
        raw: dict[str, Any] = {}
        for index, cell in enumerate(row):
            if cell is None or index >= len(header):
                continue
            title = str(header[index] or "").strip()
            text = str(cell).strip()
            if not text:
                continue
            raw[title or f"col_{index}"] = text
            lowered = title.lower()
            if lowered and any(lowered.startswith(prefix) for prefix in AGENCY_PREFIXES):
                agency[title] = text

        return IncidentTypeImportItem(
            code=value("code") or f"row-{row_no}",
            name=name or f"{category_name} (без наименования)",
            category_name=category_name,
            features=features,
            main_service=value("main_service"),
            agency_classifiers=agency,
            raw=raw,
        )


def normalize_service(name: str | None) -> str | None:
    """Название службы из классификатора → код службы оповещения."""
    if not name:
        return None
    cleaned = str(name).strip().lower()
    if not cleaned:
        return None
    if cleaned in SERVICE_ALIASES:
        return SERVICE_ALIASES[cleaned]
    for alias, code in SERVICE_ALIASES.items():
        if alias in cleaned:
            return code
    return cleaned[:32]
