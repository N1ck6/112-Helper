from __future__ import annotations

import io
import json
from xml.etree import ElementTree as ET

from httpx import AsyncClient


async def _run_attestation(
    client: AsyncClient, teacher_headers, student_headers, student_id, categories, *, solve: bool
) -> str:
    """Проводит аттестационное занятие из одной карточки и возвращает его id."""
    created = await client.post(
        "/api/v1/lessons",
        headers=teacher_headers,
        json={
            "title": "Аттестация операторов ДДС",
            "mode": "card_action",
            "purpose": "attestation",
            "passing_score": 60,
            "student_ids": [student_id],
            "category_ids": [categories["fire_apartment"]],
            "max_cards": 1,
        },
    )
    assert created.status_code == 201, created.text
    lesson_id = created.json()["id"]
    assert created.json()["purpose"] == "attestation"
    assert created.json()["passing_score"] == 60

    await client.post(f"/api/v1/lessons/{lesson_id}/start", headers=teacher_headers)
    await client.post(f"/api/v1/training/lessons/{lesson_id}/join", headers=student_headers)
    issued = await client.post(
        f"/api/v1/training/lessons/{lesson_id}/next-card", headers=student_headers
    )
    attempt_id = issued.json()["attempt"]["id"]

    if solve:
        chain = [
            ("accepted", "Карточка принята, наряд формируется"),
            ("response_started", "Подразделение направлено к месту происшествия"),
            ("arrived", "Прибытие подразделения на место"),
            ("work_in_progress", "Проводится тушение и эвакуация жильцов"),
            ("work_completed", "Работы завершены, информация передана в диспетчерскую"),
        ]
        for status_code, comment in chain:
            response = await client.post(
                f"/api/v1/training/attempts/{attempt_id}/response-status",
                headers=student_headers,
                json={"status": status_code, "comment": comment},
            )
            assert response.status_code == 201, response.text

    finished = await client.post(
        f"/api/v1/lessons/{lesson_id}/finish",
        headers=teacher_headers,
        json={"reason": "Аттестация завершена"},
    )
    assert finished.status_code == 200, finished.text
    return lesson_id


async def test_attestation_gives_verdict_and_protocol(
    client: AsyncClient, teacher_headers, student_headers, student_id, categories
) -> None:
    """Аттестация подводит итог, выставляет зачёт и печатается протоколом."""
    lesson_id = await _run_attestation(
        client, teacher_headers, student_headers, student_id, categories, solve=True
    )

    lesson = await client.get(f"/api/v1/lessons/{lesson_id}", headers=teacher_headers)
    participant = lesson.json()["participants"][0]
    assert participant["final_score"] is not None
    assert participant["is_passed"] is True, participant

    protocol = await client.post(
        "/api/v1/reports",
        headers=teacher_headers,
        json={"type": "attestation", "format": "pdf", "lesson_id": lesson_id},
    )
    assert protocol.status_code == 201, protocol.text
    body = protocol.json()
    assert body["status"] == "ready", body.get("error")
    assert body["data"]["passing_score"] == 60
    assert body["data"]["rows"][0]["verdict"] == "зачёт"
    assert body["generation_ms"] < 30_000, "Норматив формирования отчёта — 30 секунд"

    downloaded = await client.get(
        f"/api/v1/reports/{body['id']}/download", headers=teacher_headers
    )
    assert downloaded.status_code == 200
    assert downloaded.content[:4] == b"%PDF"


async def test_certificate_requires_passed_attestation(
    client: AsyncClient, teacher_headers, student_headers, student_id, categories
) -> None:
    """Без зачёта сертификат не выдаётся: он подтверждает прохождение программы."""
    await _run_attestation(
        client, teacher_headers, student_headers, student_id, categories, solve=False
    )

    response = await client.post(
        "/api/v1/certificates",
        headers=teacher_headers,
        json={"student_id": student_id, "program_name": "Подготовка оператора ДДС", "hours": 72},
    )
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "attestation_not_passed"


async def test_report_exports_to_xlsx_and_xml(
    client: AsyncClient, teacher_headers, student_headers, student_id, categories
) -> None:
    """Отчёт выгружается в Excel и в XML для legacy-интеграций."""
    lesson_id = await _run_attestation(
        client, teacher_headers, student_headers, student_id, categories, solve=True
    )

    for fmt, check in (("xlsx", b"PK"), ("xml", b"<?xml")):
        created = await client.post(
            "/api/v1/reports",
            headers=teacher_headers,
            json={"type": "lesson", "format": fmt, "lesson_id": lesson_id},
        )
        assert created.status_code == 201, created.text
        assert created.json()["status"] == "ready", created.json().get("error")

        downloaded = await client.get(
            f"/api/v1/reports/{created.json()['id']}/download", headers=teacher_headers
        )
        assert downloaded.status_code == 200
        assert downloaded.content.startswith(check), fmt

    #: Excel-файл должен реально открываться и содержать оба листа.
    from openpyxl import load_workbook

    created = await client.post(
        "/api/v1/reports",
        headers=teacher_headers,
        json={"type": "attestation", "format": "xlsx", "lesson_id": lesson_id},
    )
    downloaded = await client.get(
        f"/api/v1/reports/{created.json()['id']}/download", headers=teacher_headers
    )
    workbook = load_workbook(io.BytesIO(downloaded.content))
    assert workbook.sheetnames == ["Данные", "Сводка"]
    assert workbook["Данные"].max_row >= 2


async def test_cards_exchange_xml_roundtrip(client: AsyncClient, teacher_headers, categories) -> None:
    """Карточки выгружаются в XML и загружаются обратно пакетом (п.2.9)."""
    export = await client.get(
        "/api/v1/cards/export",
        headers=teacher_headers,
        params={"format": "xml", "limit": 2, "with_expected": True},
    )
    assert export.status_code == 200, export.text
    assert export.headers["content-type"].startswith("application/xml")

    root = ET.fromstring(export.text)
    assert root.tag == "incident-cards"
    cards = root.findall("card")
    assert cards, "В учебной базе должны быть готовые карточки"
    codes = {field.get("code") for field in cards[0].findall("fields/field")}
    assert {"address_street", "incident_class"} & codes or cards[0].find("expected-fields") is not None

    #: Меняем номера, чтобы импорт не посчитал карточки дублями.
    for index, card in enumerate(cards, start=1):
        card.set("no", f"IMP-{index:04d}")
    payload = ET.tostring(root, encoding="unicode")

    imported = await client.post(
        "/api/v1/cards/import",
        headers=teacher_headers,
        params={"category_id": categories["gas"]},
        files={"file": ("cards.xml", payload.encode("utf-8"), "application/xml")},
    )
    assert imported.status_code == 201, imported.text
    body = imported.json()
    assert body["imported"] == len(cards)
    assert body["skipped"] == 0

    #: Повторная загрузка того же пакета ничего не перезаписывает.
    again = await client.post(
        "/api/v1/cards/import",
        headers=teacher_headers,
        files={"file": ("cards.xml", payload.encode("utf-8"), "application/xml")},
    )
    assert again.json()["imported"] == 0
    assert again.json()["skipped"] == len(cards)
    assert "уже есть" in again.json()["warnings"][0]


async def test_cards_import_accepts_json_batch(client: AsyncClient, teacher_headers, categories) -> None:
    """Пакетный импорт принимает и JSON — формат обмена современных систем."""
    batch = {
        "version": "1.0",
        "cards": [
            {
                "card_no": "JSON-0001",
                "title": "Задымление на лестничной клетке",
                "difficulty": "basic",
                "fields": {
                    "aon_phone": "+7 (495) 000-11-11",
                    "applicant_name": "Соколов Пётр Ильич",
                    "applicant_role": "очевидец",
                    "address_region": "г. Москва",
                    "address_district": "САО",
                    "address_street": "улица Флотская",
                    "address_house": "17",
                    "survey_signs": ["дым", "подъезд"],
                    "incident_class": "Задымление",
                    "has_victims": False,
                    "description": "Задымление на лестничной клетке между третьим и четвёртым этажом",
                },
            }
        ],
    }
    response = await client.post(
        "/api/v1/cards/import",
        headers=teacher_headers,
        params={"category_id": categories["fire"]},
        files={"file": ("batch.json", json.dumps(batch).encode("utf-8"), "application/json")},
    )
    assert response.status_code == 201, response.text
    assert response.json()["imported"] == 1

    card_id = response.json()["card_ids"][0]
    card = await client.get(f"/api/v1/cards/{card_id}", headers=teacher_headers)
    #: Список оповещения пересчитан по ЕКП категории, а не взят из файла.
    services = {item["code"] for item in card.json()["notification_list"]}
    assert "101" in services
